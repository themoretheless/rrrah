"""Verify all pinned source/docs/notebooks/data/media and all original tests plus native boundaries."""
from pathlib import Path
import json,hashlib,csv,zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
from PIL import Image
R=Path(__file__).resolve().parents[1];W=R/'.dedup-work/deduplipy99-20261016';S=W/'source';D=R/'docs/research/dedup-500';h=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def emit(p,d):assert not p.exists();p.write_text(json.dumps(d,indent=2)+'\n')
def obj(k,b):return hashlib.sha1(k.encode()+b' '+str(len(b)).encode()+b'\0'+b).hexdigest()
g=json.loads((W/'github-commit.json').read_text());t=json.loads((D/'evidence/fritshermans--deduplipy/tree.json').read_text());assert g['tree']['sha']==t['sha'] and not t['truncated'];entries=t['tree'];blobs=[]
for e in entries:
 if e['type']!='blob':continue
 p=S/e['path'];b=p.read_bytes();assert len(b)==e['size'] and obj('blob',b)==e['sha'];row={**e,'sha256':h(p)}
 if p.suffix not in ['.png','.xlsx']:row['text_content']=b.decode()
 blobs.append(row)
expected={'':t['sha'],**{e['path']:e['sha'] for e in entries if e['type']=='tree'}}
for dirname,want in expected.items():
 children=[e for e in entries if ('' if str(Path(e['path']).parent)=='.' else str(Path(e['path']).parent))==dirname];children.sort(key=lambda e:(Path(e['path']).name+('/' if e['type']=='tree' else '')).encode());b=b''.join(('40000' if e['type']=='tree' else e['mode']).encode()+b' '+Path(e['path']).name.encode()+b'\0'+bytes.fromhex(e['sha']) for e in children);assert obj('tree',b)==want
assert len(blobs)==62
records=json.loads((W/'data-review/records.json').read_text());rows=list(csv.DictReader((S/'deduplipy/data/voter_names.csv').open()));assert records['csv']['rows']==rows and len(rows)==1746 and all(set(r)=={'name','suburb','postcode'} and all(r.values()) for r in rows)
book=load_workbook(S/'deduplipy/data/stoxx50_extended_with_id.xlsx',read_only=True,data_only=False)
for sheet,stored in zip(book,records['xlsx']):assert sheet.title==stored['name'] and [list(r) for r in sheet.iter_rows(values_only=True)]==stored['rows'] and not stored['formulas'];assert sheet.max_row==135 and sheet.max_column==2
with zipfile.ZipFile(S/'deduplipy/data/stoxx50_extended_with_id.xlsx') as z:
 assert len(z.infolist())==len(records['xlsx_members'])
 for e in records['xlsx_members']:assert hashlib.sha256(z.read(e['path'])).hexdigest()==e['sha256']
notebooks=[];cells=0
for p in [S/'docs/source/Tutorial.ipynb',S/'notebooks/advanced_example.ipynb',S/'notebooks/simple_example.ipynb']:
 n=json.loads(p.read_text());mimes=set();output_count=0
 for c in n['cells']:
  assert c['cell_type'] in ['code','markdown'];assert not c.get('attachments')
  for out in c.get('outputs',[]):mimes.update(out.get('data',{}));output_count+=1
 assert mimes<={'text/html','text/plain'};cells+=len(n['cells']);notebooks.append({'path':str(p.relative_to(R)),'sha256':h(p),'cells':len(n['cells']),'outputs':output_count,'mimes':sorted(mimes),'scope':'Allcellsource/markdown andstoredtextoutputs read; HTMLdataframeoutputvalues visibleintext/plain also read. No notebookexecution/pickleloading orhistoricalresult reproduction.'})
assert cells==73;im=Image.open(S/'docs/source/_static/logo.png');im.load();media={'path':str((S/'docs/source/_static/logo.png').relative_to(R)),'sha256':h(S/'docs/source/_static/logo.png'),'size':list(im.size),'rgba_sha256':hashlib.sha256(im.convert('RGBA').tobytes()).hexdigest(),'viewed_full':True}
proofs=[];findings=[]
for label in ['clustering','blocking','sampling','learning']:
 p=D/f'deduplipy99-{label}-controls-20261016.json';d=json.loads(p.read_text());assert h(R/f'scripts/qualify-deduplipy99-{label}-controls-20261016.py')==d['producer_sha256'];assert d['runs'][0]==d['runs'][1]
 for key in ['source_pins','loaded_dependency_file_pins']:
  for name,want in d.get(key,{}).items():assert h(Path(name) if Path(name).is_absolute() else R/name)==want
 proofs.append({'path':str(p.relative_to(R)),'sha256':h(p),'scope':d['scope']});findings.extend(d['findings'])
p=D/'deduplipy99-full-tests-20261016.json';d=json.loads(p.read_text());assert h(R/'scripts/qualify-deduplipy99-full-tests-20261016.py')==d['producer_sha256'];assert d['runs'][0]['cases']==d['runs'][1]['cases'] and len(d['runs'][0]['cases'])==49
for run in d['runs']:
 assert run['exit_code']==0 and h(R/run['log'])==run['log_sha256'] and h(R/run['junit'])==run['junit_sha256'];xml=ET.parse(R/run['junit']).getroot();assert len(list(xml.iter('testcase')))==49 and not list(xml.iter('failure')) and not list(xml.iter('error'))
for name,want in d['artifact_pins'].items():assert h(R/name)==want
assert d['predict_controls'][0]==d['predict_controls'][1] and len(d['predict_controls'][0])==4;findings.extend(d['findings']);proofs.append({'path':str(p.relative_to(R)),'sha256':h(p),'scope':d['scope']})
a=D/'deduplipy99-complete-source-pins-20261016.json';emit(a,{'repository':'fritshermans/deduplipy','commit':g['sha'],'canonical_trees':expected,'blobs':blobs,'notebooks':notebooks,'media':media,'data_review':{'path':str((W/'data-review/records.json').relative_to(R)),'sha256':h(W/'data-review/records.json'),'CSV_records':1746,'XLSX_data_rows':134,'scope':'Allrecords parsed/type/field/member integrityreview; everyXLSXcellvalue read. CSVnames notindividuallyidentityvalidated. No formulas/macros/externalworkbookexecution.'}})
emit(D/'deduplipy99-complete-source-review-20261016.json',{'status':'verified_complete_pinned_source_data_media_all_original_tests_and_native_boundaries','whole_repo_complete':True,'whole_implementation_complete':True,'repository':'fritshermans/deduplipy','commit':g['sha'],'blobs':62,'canonical_trees':len(expected),'acquisition':str(a.relative_to(R)),'acquisition_sha256':h(a),'auditor_sha256':h(Path(__file__)),'native_proofs':proofs,'upstream_tests_run':True,'upstream_test_scope':'All49originaltests twice98passes inresolvedMacARM/Python3.12deps; seed1234forrepeatability. Earlierselectedsubsets overlap samecases, not98additionaluniquetests. No originalrelease/historicalenvironment guarantee.','additional_native_controls':90,'native_scope':'45selectedboundarycases twice acrossclustering/blocking/sampling/learning/fullpredict. Realnumeric/MLdependencies; onlyinteractivebuiltininput replaced inlearningstage. Fullpredict usesactuallearnertrainedsyntheticfeatures+actualBlockingfit; no body/algorithmadapter.','media_scope':'Trackedlogo viewedfull; all73cells/outputs of3notebooks read, text/HTMLdataframes only,noembeddedimage/video. AllCSV/XLSXrecords structurallyreviewed andXLSXallcells read. Externaldynamicbadges/liveReadTheDocslogo notfetched; no linkedsite/fullversioncrawl claim.','license':'Complete MIT LICENSE read. No productioncodecopied.','findings':findings,'decision':'Reference learnedcandidate selection/featurepipeline only withprovenrecall/index/objectidentity/negative-edge preservation andempty-input contracts. Entityresolution/centroidimputation isnotimage/file/pixel identity. No productionadoption orspeedclaim.','limits':['Fullsource reading andalloriginaltests donotprove realcorpus/entity/image precision/recall, million-scale/hardmemory/thread/platform/resources/performance.','Historicalnotebookoutputs read, notexecuted/reproduced. CSVrecordidentity notmanuallyvalidated; no privateuserdata/network/pickle/sourceuploads.','Knownnativecounterexamples remain; originaltests passing isnotuniversalcorrectness orcurrentrootqualification.']});print(json.dumps({'status':'verified','blobs':62,'trees':len(expected),'notebook_cells':73,'upstream_passes':98,'additional_controls':90}))
